"""어드민 캠페인 로그 엑셀 추출 (2026-09-23 JDH).

시트 3장, 모두 맨 왼쪽 두 열은 진행한 회원의 로그인 ID 와 회사명이다.
  캠페인   — 1행 = 캠페인 1건. 스토어 상품번호 · nvMid · 쿠팡 상품ID · 플레이스ID 를 각각 다른 열에.
  일별 로그 — 1행 = 캠페인 × 날짜. 예정 수량과 어드민이 기록한 실작업량·순위.
  주별 정산 — 주(월요일 시작) × 회원 × 채널 × 매체. 이용 일수·수량·금액, 환불은 별도 행에 −로.

일별·주별은 "이용"이 실제로 일어난 날만 센다: 상태가 running/done/stopped 이고, 시작일 이후 ~
min(종료일, 오늘, 중단일) 까지. 검수·승인 대기·반려·취소는 이용 행이 없고 환불 행만 나올 수 있다.

nvMid 는 URL 에 없다(스마트스토어 주소의 숫자는 스토어 상품번호). 위저드 미리보기(순위 서버)가
돌려준 값을 campaigns.nv_mid 에 저장해 두고 여기서 읽는다. 없으면 빈칸.
"""
import io
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from urllib.parse import parse_qs, urlparse

from ..constants import CHANNEL_LABEL, PAY_METHOD_LABEL, STATUS_LABEL

USED_STATUSES = ("running", "done", "stopped")

_STORE_NO = re.compile(r"(?:smartstore|brand|m\.smartstore|m\.brand)\.naver\.com/[^/?#]+/products/(\d+)")
_NVMID_PATH = re.compile(r"shopping\.naver\.com/(?:catalog|product)/(\d+)")
_COUPANG = re.compile(r"coupang\.com/vp/products/(\d+)")
_PLACE_PATTERNS = (
    re.compile(r"place\.naver\.com/[a-z]+/(\d{5,})"),          # m.place / pcmap.place: /restaurant/123/home
    re.compile(r"map\.naver\.com/p/(?:entry/)?place/(\d{5,})"),
    re.compile(r"map\.naver\.com/v5/entry/place/(\d{5,})"),
    re.compile(r"[?&]placeId=(\d{5,})"),
    re.compile(r"[?&]id=(\d{5,})"),
)


def store_product_no(url):
    """스마트스토어·브랜드스토어 주소의 상품번호 (/products/NNN)."""
    m = _STORE_NO.search(url or "")
    return m.group(1) if m else None


def nv_mid_from_url(url):
    """네이버 쇼핑 카탈로그/상품 주소 또는 ?nvMid= 에 들어 있는 nvMid. 스토어 주소에는 없다."""
    if not url:
        return None
    q = parse_qs(urlparse(url).query)
    if q.get("nvMid") and q["nvMid"][0].isdigit():
        return q["nvMid"][0]
    m = _NVMID_PATH.search(url)
    return m.group(1) if m else None


def coupang_product_id(url):
    m = _COUPANG.search(url or "")
    return m.group(1) if m else None


def place_id(url):
    """플레이스 URL에서 플레이스 ID. naver.me 단축 주소는 풀 수 없어 None."""
    if not url:
        return None
    for pat in _PLACE_PATTERNS:
        m = pat.search(url)
        if m:
            return m.group(1)
    return None


def ids_of(c):
    """캠페인 한 건의 식별자 4종. 채널에 맞는 것만 채운다."""
    url = c.get("target_url")
    out = {"store_no": None, "nv_mid": None, "coupang": None, "place": None}
    if c["channel"] == "store":
        out["store_no"] = store_product_no(url)
        out["nv_mid"] = c.get("nv_mid") or nv_mid_from_url(url)
    elif c["channel"] == "coupang":
        out["coupang"] = coupang_product_id(url)
    elif c["channel"] == "place":
        out["place"] = place_id(url)
    return out


def login_id(c):
    """캠페인 행의 회원 로그인 아이디 — 표기 규칙은 models/user.login_id 한 곳에서."""
    from ..models.user import login_id as fmt
    return fmt(email=c.get("user_email"), kakao_id=c.get("user_kakao"), user_id=c.get("user_id"))


def _week_start(d):
    return d - timedelta(days=d.weekday())


def _used_range(c, today):
    """(first, last) 실제 이용 구간. 없으면 None."""
    if c["status"] not in USED_STATUSES or not c.get("start_date"):
        return None
    first, last = c["start_date"], c["end_date"] or c["start_date"]
    if c["status"] == "stopped" and c.get("updated_at"):
        u = c["updated_at"].date() if isinstance(c["updated_at"], datetime) else c["updated_at"]
        last = min(last, u)
    last = min(last, today)
    return (first, last) if first <= last else None


def _days(first, last):
    d = first
    while d <= last:
        yield d
        d += timedelta(days=1)


def build_workbook(rows, daily, refunds, date_from, date_to, today=None):
    """rows: campaign_model.admin_all 결과. daily: campaign_model.daily_records. refunds: credit_model.campaign_refunds."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    today = today or date.today()
    wb = Workbook()
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EEECFA")
    neg_font = Font(color="C0392B")

    def sheet(title, headers, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "C2"     # 회원 ID·회사명 두 열은 가로 스크롤에도 고정
        return ws

    def who(c):
        return [login_id(c), c.get("user_biz")]

    # ---------------------------------------------------------------- 캠페인
    ws = sheet("캠페인",
               ["회원 ID", "회사명", "주문번호", "상태", "채널", "닉네임", "상품/업체명", "키워드", "매체",
                "스토어 상품번호", "nvMid", "쿠팡 상품ID", "플레이스ID", "링크",
                "시작", "종료", "일수", "일 수량", "총 수량", "단가", "결제 금액", "환불(−)", "순매출",
                "결제수단", "시작 순위", "현재 순위", "결제일", "등록일"],
               [22, 16, 16, 9, 10, 12, 22, 14, 10, 16, 16, 16, 14, 44, 11, 11, 6, 8, 9, 7, 12, 11, 12, 10, 8, 8, 17, 17])
    for c in rows:
        ids = ids_of(c)
        days = (c["end_date"] - c["start_date"]).days + 1 if c.get("start_date") and c.get("end_date") else None
        refund = c.get("refund_amount") or 0
        ws.append(who(c) + [
            c["order_no"], STATUS_LABEL.get(c["status"], c["status"]), CHANNEL_LABEL.get(c["channel"], c["channel"]),
            c.get("nickname"), c.get("product_name") or c.get("biz_name"), c.get("main_keyword"), c.get("media_name"),
            ids["store_no"], ids["nv_mid"], ids["coupang"], ids["place"], c["target_url"],
            c.get("start_date"), c.get("end_date"), days, c.get("daily_qty"), c.get("total_qty"), c.get("unit_price"),
            c.get("paid_amount"), -refund if refund else 0, (c.get("paid_amount") or 0) - refund,
            PAY_METHOD_LABEL.get(c.get("pay_method"), c.get("pay_method")), c.get("rank_start"), c.get("rank_now"),
            c.get("paid_at"), c.get("created_at")])
        if refund:
            ws.cell(row=ws.max_row, column=22).font = neg_font

    # ---------------------------------------------------------------- 일별 로그
    ws = sheet("일별 로그",
               ["회원 ID", "회사명", "날짜", "요일", "주문번호", "채널", "매체", "상품/업체명", "키워드",
                "예정 수량", "기록 수량", "순위", "단가", "예정 금액"],
               [22, 16, 11, 5, 16, 10, 10, 22, 14, 9, 9, 6, 7, 11])
    yoil = "월화수목금토일"
    weekly = defaultdict(lambda: {"campaigns": set(), "days": 0, "qty": 0, "done": 0, "done_seen": False, "amount": 0})
    for c in rows:
        rng = _used_range(c, today)
        if not rng:
            continue
        first, last = max(rng[0], date_from), min(rng[1], date_to)
        if first > last:
            continue
        for d in _days(first, last):
            rec = daily.get((c["id"], d))
            done = rec["done_qty"] if rec and rec.get("done_qty") is not None else None
            ws.append(who(c) + [
                d, yoil[d.weekday()], c["order_no"], CHANNEL_LABEL.get(c["channel"], c["channel"]), c.get("media_name"),
                c.get("product_name") or c.get("biz_name"), c.get("main_keyword"),
                c["daily_qty"], done, rec["rank"] if rec else None, c["unit_price"], c["daily_qty"] * c["unit_price"]])
            w = weekly[(_week_start(d), c["user_id"], c["channel"], c["media_id"])]
            w["campaigns"].add(c["order_no"])
            w["days"] += 1
            w["qty"] += c["daily_qty"]
            w["amount"] += c["daily_qty"] * c["unit_price"]
            w.setdefault("who", who(c))
            w.setdefault("media_name", c.get("media_name"))
            if done is not None:
                w["done"] += done
                w["done_seen"] = True

    # ---------------------------------------------------------------- 주별 정산
    ws = sheet("주별 정산",
               ["회원 ID", "회사명", "주 시작(월)", "주 종료(일)", "채널", "매체", "항목", "캠페인 수", "캠페인",
                "이용 일수", "이용 수량", "기록 수량", "금액", "비고"],
               [22, 16, 12, 12, 10, 10, 7, 9, 30, 9, 10, 10, 13, 30])
    by_id = {c["id"]: c for c in rows}
    lines = []
    for (wk, uid, ch, mid), w in weekly.items():
        lines.append((wk, w["who"][0] or "", ch, w["media_name"] or "", 0,
                      w["who"] + [wk, wk + timedelta(days=6), CHANNEL_LABEL.get(ch, ch), w["media_name"], "이용",
                                  len(w["campaigns"]), ", ".join(sorted(w["campaigns"])),
                                  w["days"], w["qty"], w["done"] if w["done_seen"] else None, w["amount"], None]))
    for r in refunds:
        c = by_id.get(r["campaign_id"])
        if not c:
            continue
        rd = r["created_at"].date() if isinstance(r["created_at"], datetime) else r["created_at"]
        if not (date_from <= rd <= date_to):
            continue
        wk = _week_start(rd)
        per_day = (c["daily_qty"] or 0) * (c["unit_price"] or 0)
        days_eq = -round(r["amount"] / per_day) if per_day else None
        qty_eq = -(r["amount"] // c["unit_price"]) if c.get("unit_price") else None
        lines.append((wk, login_id(c), c["channel"], c.get("media_name") or "", 1,
                      who(c) + [wk, wk + timedelta(days=6), CHANNEL_LABEL.get(c["channel"], c["channel"]),
                                c.get("media_name"), "환불", 1, c["order_no"], days_eq, qty_eq, None, -r["amount"],
                                f"{rd:%Y.%m.%d} {r.get('memo') or ''}".strip()]))
    lines.sort(key=lambda x: (x[0], x[1], x[2], x[3], x[4]))
    for _, _, _, _, kind, vals in lines:
        ws.append(vals)
        if kind == 1:
            for col in (10, 11, 13):
                ws.cell(row=ws.max_row, column=col).font = neg_font
    ws.append([])
    ws.append([f"기간 {date_from:%Y.%m.%d} ~ {date_to:%Y.%m.%d} · 이용 = 구동 중·완료·중단 캠페인의 실제 구동일 기준(오늘까지) · "
               "환불은 크레딧 원장의 환불 시점 주에 − 로 기록 · 기록 수량은 어드민이 입력한 일별 작업량 합(없으면 빈칸)"])

    wb.remove(wb["Sheet"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
