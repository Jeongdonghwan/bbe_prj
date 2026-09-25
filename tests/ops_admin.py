"""운영 점검 — 어드민 화면. 승인·반려·순위·크레딧·콘텐츠까지 실제로 처리해 본다."""
import io
import os
import sys
from datetime import date, timedelta

os.chdir(r"C:\bbe_prj")
sys.path.insert(0, r"C:\bbe_prj")

from app import create_app  # noqa: E402
from app.db import execute, query, query_one  # noqa: E402

FAIL = []


def check(cond, name, extra=""):
    print(("PASS " if cond else "FAIL ") + name + (("  → " + str(extra)) if (extra and not cond) else ""))
    if not cond:
        FAIL.append(name)


def body(r):
    return r.get_data(as_text=True)


def main():
    app = create_app()
    app.config["TESTING"] = True
    admin = app.test_client()
    user = app.test_client()
    with app.app_context():
        aid = query_one("SELECT id FROM users WHERE role='admin' LIMIT 1")["id"]
        uid = query_one("SELECT id FROM users WHERE role<>'admin' AND status='active' ORDER BY id LIMIT 1")["id"]
        execute("UPDATE users SET credit_balance=5000000 WHERE id=%s", [uid])
        m = query_one("SELECT id, unit_price FROM media WHERE channel='store' AND is_active=1 ORDER BY sort LIMIT 1")
        from app.services import campaign_service
        start = campaign_service.earliest_start()
    with admin.session_transaction() as s:
        s["uid"] = aid
    with user.session_transaction() as s:
        s["uid"] = uid

    # 점검용 캠페인 2건 (승인용 / 반려용)
    total = m["unit_price"] * 300 * 10
    made = []
    for i in (1, 2):
        user.post("/campaign/store/new", data={
            "media_id": m["id"], "days": "10", "daily_qty": "300", "product_name": f"어드민점검 {i}",
            "target_url": f"https://smartstore.naver.com/ops/products/9000{i}", "main_keyword": f"어드민점검키워드{i}",
            "start_date": start.isoformat(), "client_total": str(total)}, follow_redirects=True)
    with app.app_context():
        made = query("SELECT * FROM campaigns WHERE biz_name LIKE '어드민점검%' ORDER BY id")
    check(len(made) == 2, "점검용 캠페인 2건 생성", len(made))
    ok_c, no_c = made[0], made[1]

    print("\n=== 1. 운영 현황 ===")
    h = body(admin.get("/admin"))
    check("검수" in h or "대기" in h, "대시보드 위젯")

    print("\n=== 2. 주문 관리 — 승인 → 순위 → 완료 ===")
    h = body(admin.get("/admin/orders?status=review"))
    check("어드민점검 1" in h, "검수 대기 목록에 노출")
    # 승인 = 곧장 정상(running). 중간에 "구동 시작"을 누르는 단계는 없다.
    r = admin.post(f"/admin/orders/{ok_c['id']}/action", data={"action": "status", "status": "running"}, follow_redirects=True)
    with app.app_context():
        st = query_one("SELECT status FROM campaigns WHERE id=%s", [ok_c["id"]])["status"]
    check(st == "running", "승인 → 정상", st)
    r = admin.post(f"/admin/orders/{ok_c['id']}/rank", data={"rank": "17", "done_qty": "300"}, follow_redirects=True)
    with app.app_context():
        c2 = query_one("SELECT rank_start, rank_now FROM campaigns WHERE id=%s", [ok_c["id"]])
        d = query_one("SELECT `rank`, done_qty FROM campaign_daily WHERE campaign_id=%s AND date=CURDATE()", [ok_c["id"]])
    check(c2["rank_now"] == 17 and c2["rank_start"] == 17 and d and d["rank"] == 17, "순위 입력", dict(c2))
    with app.app_context():
        nt = query_one("SELECT COUNT(*) n FROM notifications WHERE user_id=%s", [uid])["n"]
    check(nt > 0, "사용자 알림 생성", nt)

    print("\n=== 3. 주문 관리 — 반려 + 크레딧 환불 ===")
    with app.app_context():
        bal0 = query_one("SELECT credit_balance FROM users WHERE id=%s", [uid])["credit_balance"]
    r = admin.post(f"/admin/orders/{no_c['id']}/action", data={"action": "status", "status": "rejected", "reason": "점검 반려"},
                   follow_redirects=True)
    with app.app_context():
        c3 = query_one("SELECT status, refund_amount FROM campaigns WHERE id=%s", [no_c["id"]])
        bal1 = query_one("SELECT credit_balance FROM users WHERE id=%s", [uid])["credit_balance"]
    check(c3["status"] == "rejected", "반려", c3["status"])
    check(bal1 - bal0 == total, "반려 시 전액 환불", f"{bal0}→{bal1} (기대 +{total})")
    r = admin.post(f"/admin/orders/{no_c['id']}/action", data={"action": "status", "status": "rejected"}, follow_redirects=True)
    check("사유" in body(r) or "허용되지" in body(r), "사유 없는 반려 거절")

    print("\n=== 4. 주문 관리 — 메모 · 일괄 · 엑셀 ===")
    r = admin.post(f"/admin/orders/{ok_c['id']}/action", data={"action": "memo", "memo": "점검 메모"}, follow_redirects=True)
    with app.app_context():
        mm = query_one("SELECT admin_memo FROM campaigns WHERE id=%s", [ok_c["id"]])["admin_memo"]
    check(mm == "점검 메모", "관리자 메모", mm)
    r = admin.get("/admin/orders/export?from=%s" % (date.today() - timedelta(days=60)).isoformat())
    check(r.status_code == 200 and len(r.data) > 5000, "엑셀 내보내기", len(r.data))
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    check(wb.sheetnames == ["캠페인", "일별 로그", "주별 정산"], "엑셀 3시트", wb.sheetnames)

    print("\n=== 5. 순위 일괄 업로드 ===")
    csv = "주문번호,날짜,순위,작업량\n%s,%s,9,300\n" % (ok_c["order_no"], date.today().isoformat())
    r = admin.post("/admin/orders/rank-upload",
                   data={"file": (io.BytesIO(csv.encode("utf-8-sig")), "ranks.csv")},
                   content_type="multipart/form-data", follow_redirects=True)
    with app.app_context():
        d = query_one("SELECT `rank` FROM campaign_daily WHERE campaign_id=%s AND date=CURDATE()", [ok_c["id"]])
    check(d and d["rank"] == 9, "CSV 순위 업로드", (d or {}).get("rank"))

    print("\n=== 6. 크레딧 관리 ===")
    user.post("/credit/charge", data={"amount": "200000", "depositor": "점검", "tax_invoice": "0"}, follow_redirects=True)
    with app.app_context():
        req = query_one("SELECT * FROM charge_requests WHERE user_id=%s ORDER BY id DESC LIMIT 1", [uid])
        bal0 = query_one("SELECT credit_balance FROM users WHERE id=%s", [uid])["credit_balance"]
    h = body(admin.get("/admin/credits"))
    check("점검" in h, "충전 요청 목록")
    r = admin.post(f"/admin/credits/{req['id']}/approve", follow_redirects=True)
    with app.app_context():
        st = query_one("SELECT status FROM charge_requests WHERE id=%s", [req["id"]])["status"]
        bal1 = query_one("SELECT credit_balance FROM users WHERE id=%s", [uid])["credit_balance"]
    check(st == "approved" and bal1 - bal0 == 200000, "충전 승인 → 잔액 반영", f"{st} {bal0}→{bal1}")
    r = admin.post("/admin/credits/adjust", data={"user_id": uid, "amount": "-50000", "memo": "점검 차감"}, follow_redirects=True)
    with app.app_context():
        bal2 = query_one("SELECT credit_balance FROM users WHERE id=%s", [uid])["credit_balance"]
    check(bal1 - bal2 == 50000, "직접 차감", f"{bal1}→{bal2}")

    print("\n=== 7. 매체사 관리 ===")
    h = body(admin.get("/admin/media?channel=store"))
    check("이번 주 운영팀 추천" in h and "오늘의 인기" in h, "매체 화면")
    with app.app_context():
        ids = {x["name"]: x["id"] for x in query("SELECT id, name FROM media WHERE channel='store' AND is_active=1")}
    r = admin.post("/admin/media/weekly", data={"channel": "store", "rank1": ids["타이탄"], "rank2": ids["싱크"],
                                                "rank3": ids["말론"]}, follow_redirects=True)
    check("저장했습니다" in body(r), "주간 추천 저장")
    r = admin.post("/admin/media/picks", data={"channel": "store", "type_ids": [ids["타이탄"], ids["사도"]]}, follow_redirects=True)
    with app.app_context():
        from app.models import daily_pick
        picks = daily_pick.get("store", date.today())
    check(len(picks) == 2, "오늘의 인기 저장", picks)
    r = admin.post(f"/admin/media/{ids['사도']}/toggle", follow_redirects=True)
    with app.app_context():
        act = query_one("SELECT is_active FROM media WHERE id=%s", [ids["사도"]])["is_active"]
    check(act == 0, "매체 노출 토글", act)
    admin.post(f"/admin/media/{ids['사도']}/toggle", follow_redirects=True)

    print("\n=== 8. 콘텐츠 (공지·정보) ===")
    r = admin.post("/admin/content/save", data={"board": "notice", "category": "공지", "title": "점검 공지",
                                                "body": "<p>점검 본문</p>", "status": "published"}, follow_redirects=True)
    with app.app_context():
        ct = query_one("SELECT * FROM contents WHERE title='점검 공지' ORDER BY id DESC LIMIT 1")
    check(ct is not None, "공지 작성")
    if ct:
        h = body(app.test_client().get("/notice"))
        check("점검 공지" in h, "사용자 공지 목록에 즉시 반영")
        admin.post(f"/admin/content/{ct['id']}/pin", follow_redirects=True)
        with app.app_context():
            pinned = query_one("SELECT is_pinned FROM contents WHERE id=%s", [ct["id"]])["is_pinned"]
        check(pinned == 1, "고정")
        admin.post(f"/admin/content/{ct['id']}/delete", follow_redirects=True)
        with app.app_context():
            gone = query_one("SELECT id FROM contents WHERE id=%s", [ct["id"]])
        check(gone is None, "삭제")

    print("\n=== 9. 배너 ===")
    # 신규 배너는 이미지가 필수 (이미지 없는 배너는 화면에 빈칸으로 나가므로 올바른 제약)
    r = admin.post("/admin/banners/save", data={"title": "점검 배너", "link": "/notice", "zone": "grid",
                                                "sort": "9", "is_active": "1"}, follow_redirects=True)
    check("이미지 파일을 선택" in body(r), "이미지 없는 배너 거절")
    from PIL import Image as _Im
    _buf = io.BytesIO(); _Im.new("RGB", (1300, 400), (120, 100, 220)).save(_buf, "PNG"); _buf.seek(0)
    r = admin.post("/admin/banners/save", data={"title": "점검 배너", "link": "/notice", "zone": "grid",
                                                "sort": "-1", "is_active": "1",
                                                "image": (_buf, "ops.png")},
                   content_type="multipart/form-data", follow_redirects=True)
    with app.app_context():
        bn = query_one("SELECT * FROM banners WHERE title='점검 배너' ORDER BY id DESC LIMIT 1")
    check(bn is not None, "배너 등록", body(r)[-200:] if bn is None else "")
    if bn:
        check(bool(bn["image_url"]), "배너 이미지 저장", bn["image_url"])
        h = body(app.test_client().get("/"))
        check(bn["image_url"] in h, "사용자 대시보드에 배너 노출")
        admin.post(f"/admin/banners/{bn['id']}/toggle", follow_redirects=True)
        with app.app_context():
            a2 = query_one("SELECT is_active FROM banners WHERE id=%s", [bn["id"]])["is_active"]
        check(a2 == 0, "배너 토글", a2)
        admin.post(f"/admin/banners/{bn['id']}/delete", follow_redirects=True)
    r = admin.post("/admin/banners/strip", data={"strip_on": "1", "strip_text": "점검 띠배너", "strip_link": "",
                                                 "strip_bg": "#2563EB"}, follow_redirects=True)
    h = body(app.test_client().get("/"))
    check("점검 띠배너" in h, "띠배너 사용자 화면 반영")
    admin.post("/admin/banners/strip", data={"strip_text": "테스트 띠배너 문구입니다", "strip_bg": "#2563EB"}, follow_redirects=True)

    print("\n=== 10. 회원 ===")
    h = body(admin.get("/admin/users"))
    check("회원" in h, "회원 목록")
    check(admin.get(f"/admin/users/{uid}/drawer").status_code == 200, "회원 상세 드로어")
    r = admin.post(f"/admin/users/{uid}/grade", data={"grade": "agency"}, follow_redirects=True)
    with app.app_context():
        g2 = query_one("SELECT grade FROM users WHERE id=%s", [uid])["grade"]
    check(g2 == "agency", "등급 변경", g2)
    admin.post(f"/admin/users/{uid}/grade", data={"grade": "biz"}, follow_redirects=True)

    print("\n=== 11. 신고 · 인기 트래픽 · 대행 ===")
    check(admin.get("/admin/reports").status_code == 200, "신고 화면")
    check(admin.get("/admin/popular").status_code == 200, "인기 트래픽 관리")
    check(admin.get("/admin/agency").status_code == 200, "대행 의뢰 화면")

    print("\n=== 12. 입금 확인 화면 ===")
    h = body(admin.get("/admin/payments"))
    check("입금" in h or "결제" in h, "입금 확인 화면")
    r = admin.post("/admin/payments/settings", data={"bank_name": "국민은행", "bank_account": "123456-00-111111",
                                                     "bank_holder": "비베네", "bank_due_days": "3"}, follow_redirects=True)
    with app.app_context():
        bk = {x["k"]: x["v"] for x in query("SELECT k, v FROM settings WHERE k LIKE 'bank%'")}
    check(bk.get("bank_account") == "123456-00-111111", "입금 계좌 설정 저장", bk)
    h = body(user.get("/credit/charge"))
    check("123456-00-111111" in h, "사용자 충전 화면에 계좌 반영")

    print("\n=== 13. 관리자 로그 ===")
    with app.app_context():
        logs = query("SELECT action, COUNT(*) n FROM admin_log WHERE created_at >= NOW() - INTERVAL 5 MINUTE GROUP BY action")
    check(len(logs) >= 6, "쓰기 작업이 admin_log 에 기록", [(x["action"], x["n"]) for x in logs])

    # ---- cleanup ----
    with app.app_context():
        for c in made:
            for t in ("campaign_daily", "status_log", "payments"):
                execute(f"DELETE FROM {t} WHERE campaign_id=%s", [c["id"]])
            execute("DELETE FROM campaigns WHERE id=%s", [c["id"]])
        execute("DELETE FROM credit_ledger WHERE user_id=%s AND (memo LIKE '%%점검%%' OR memo LIKE '%%어드민%%')", [uid])
        execute("DELETE FROM charge_requests WHERE depositor='점검'")
        execute("DELETE FROM notifications WHERE user_id=%s", [uid])
        execute("DELETE FROM contents WHERE title='점검 공지'")
        for b in query("SELECT image_url FROM banners WHERE title='점검 배너'"):
            if b["image_url"]:
                fp = os.path.join("app", b["image_url"].lstrip("/"))
                if os.path.exists(fp):
                    os.remove(fp)
        execute("DELETE FROM banners WHERE title='점검 배너'")
        execute("UPDATE users SET credit_balance=500000 WHERE id=%s", [uid])
    print("\n" + ("전부 통과" if not FAIL else "실패 %d건: %s" % (len(FAIL), ", ".join(FAIL))))


if __name__ == "__main__":
    main()
